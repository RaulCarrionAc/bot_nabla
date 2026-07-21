import sqlite3
import pandas as pd
import numpy as np
import holidays
import io
import os
from pathlib import Path
from datetime import datetime
from typing import Tuple, Optional, List
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.formatting.rule import ColorScaleRule

# --- FUNCIONES DE CARGA Y CALENDARIO ---

def _calcular_tipo_dia(fechas: pd.Series) -> pd.Series:
    """
    Calcula tipo_dia (DL/DS/DF) a partir de una serie de fechas, usando la
    librería 'holidays' para festivos de Chile.
    """
    if fechas.empty:
        return pd.Series([], dtype=str)
    anios = fechas.dt.year.unique().tolist()
    feriados_cl = holidays.Chile(years=anios)

    es_feriado = fechas.apply(lambda x: x in feriados_cl)
    es_domingo = fechas.dt.weekday == 6
    es_sabado = fechas.dt.weekday == 5

    condiciones = [es_feriado | es_domingo, es_sabado]
    valores = ["DF", "DS"]
    return pd.Series(np.select(condiciones, valores, default="DL"), index=fechas.index)


def construir_calendario_mes(anio: int, mes: int, fecha_corte=None) -> pd.DataFrame:
    """
    Construye el calendario completo (Fecha, tipo_dia) de TODOS los días de un mes.
    Si se entrega fecha_corte, trunca hasta esa fecha.
    """
    inicio_mes = pd.Timestamp(year=anio, month=mes, day=1)
    fin_mes = inicio_mes + pd.offsets.MonthEnd(0)
    
    if fecha_corte is not None:
        fecha_corte = pd.Timestamp(fecha_corte)
        if fecha_corte < fin_mes:
            fin_mes = fecha_corte

    fechas = pd.Series(pd.date_range(start=inicio_mes, end=fin_mes, freq="D"))
    return pd.DataFrame({"Fecha": fechas, "tipo_dia": _calcular_tipo_dia(fechas)})


def obtener_df_conteo_db(db_path: str, operador: str, anio: int, mes: int) -> pd.DataFrame:
    """Obtiene el conteo de expediciones reales desde SQLite para un mes específico."""
    conn = sqlite3.connect(db_path)
    query = """
        SELECT 
            Fecha, 
            Servicio as servicio, 
            Sentido as sentido, 
            Periodo as periodo,
            COUNT(*) as expediciones_observadas
        FROM expediciones
        WHERE operador = ?
          AND strftime('%Y', Fecha) = ?
          AND strftime('%m', Fecha) = ?
          AND (Estado IN ('VALIDA', 'VÁLIDA', 'Valida', 'Válida', 'valida', 'válida'))
        GROUP BY Fecha, Servicio, Sentido, Periodo
    """
    mes_str = f"{mes:02d}"
    df = pd.read_sql_query(query, conn, params=(operador, str(anio), mes_str))
    conn.close()
    
    if df.empty:
        return pd.DataFrame(columns=["Fecha", "servicio", "sentido", "periodo", "expediciones_observadas", "tipo_dia"])
        
    df["Fecha"] = pd.to_datetime(df["Fecha"])
    df["sentido"] = df["sentido"].astype(str).str.strip().str.upper().str[0]
    df["servicio"] = df["servicio"].astype(str).str.strip()
    df["tipo_dia"] = _calcular_tipo_dia(df["Fecha"])
    
    if operador.lower() == "tasacop":
        df["servicio"] = df["servicio"].str.replace(r"_(I|R)$", "", regex=True)
        # Agrupar y sumar por si quedaron registros duplicados tras limpiar el sufijo
        df = df.groupby(["Fecha", "servicio", "sentido", "periodo", "tipo_dia"], as_index=False)["expediciones_observadas"].sum()
        
    df["periodo"] = pd.to_numeric(df["periodo"], errors="coerce").astype("Int64")
    
    return df


def obtener_df_conteo_historico_db(db_path: str, operador: str, fecha_limite: pd.Timestamp) -> pd.DataFrame:
    """Obtiene el conteo de expediciones para todos los periodos anteriores al mes actual."""
    conn = sqlite3.connect(db_path)
    query = """
        SELECT 
            Fecha, 
            Servicio as servicio, 
            Sentido as sentido, 
            Periodo as periodo,
            COUNT(*) as expediciones_observadas
        FROM expediciones
        WHERE operador = ?
          AND Fecha < ?
          AND (Estado IN ('VALIDA', 'VÁLIDA', 'Valida', 'Válida', 'valida', 'válida'))
        GROUP BY Fecha, Servicio, Sentido, Periodo
    """
    limite_str = fecha_limite.strftime("%Y-%m-%d")
    df = pd.read_sql_query(query, conn, params=(operador, limite_str))
    conn.close()
    
    if df.empty:
        return pd.DataFrame(columns=["Fecha", "servicio", "sentido", "periodo", "expediciones_observadas", "tipo_dia"])
        
    df["Fecha"] = pd.to_datetime(df["Fecha"])
    df["sentido"] = df["sentido"].astype(str).str.strip().str.upper().str[0]
    df["servicio"] = df["servicio"].astype(str).str.strip()
    df["tipo_dia"] = _calcular_tipo_dia(df["Fecha"])
    
    if operador.lower() == "tasacop":
        df["servicio"] = df["servicio"].str.replace(r"_(I|R)$", "", regex=True)
        # Agrupar y sumar por si quedaron registros duplicados tras limpiar el sufijo
        df = df.groupby(["Fecha", "servicio", "sentido", "periodo", "tipo_dia"], as_index=False)["expediciones_observadas"].sum()
        
    df["periodo"] = pd.to_numeric(df["periodo"], errors="coerce").astype("Int64")
    
    return df


def obtener_df_a1_db(db_path: str, operador: str, calendario: pd.DataFrame) -> pd.DataFrame:
    """Obtiene y expande las frecuencias exigidas (Anexo 1) para un calendario de fechas."""
    conn = sqlite3.connect(db_path)
    query = """
        SELECT 
            servicio, 
            sentido, 
            periodo, 
            [Tipo de Día] as tipo_dia, 
            [Tipo Demanda] as tipo_demanda, 
            [Frecuencia (buses/hr)] as EE 
        FROM anexo_1 
        WHERE operador = ?
    """
    df_a1_raw = pd.read_sql_query(query, conn, params=(operador,))
    conn.close()
    
    if df_a1_raw.empty:
        return pd.DataFrame(columns=["Fecha", "tipo_dia", "servicio", "sentido", "periodo", "tipo_demanda", "EE"])
        
    df_a1_raw["sentido"] = df_a1_raw["sentido"].astype(str).str.strip().str.upper().str[0]
    df_a1_raw["servicio"] = df_a1_raw["servicio"].astype(str).str.strip()
    df_a1_raw["periodo"] = pd.to_numeric(df_a1_raw["periodo"], errors="coerce").astype("Int64")
    
    # Mapear los nombres de tipo de día de la base de datos a los códigos de calendario ('DL', 'DS', 'DF')
    tipo_dia_map = {
        "Laboral": "DL",
        "Sábado": "DS",
        "Domingo / Festivo": "DF"
    }
    df_a1_raw["tipo_dia"] = df_a1_raw["tipo_dia"].map(tipo_dia_map)
    
    calendario_uniq = calendario[["Fecha", "tipo_dia"]].drop_duplicates()
    df_base_exigida = pd.merge(calendario_uniq, df_a1_raw, on="tipo_dia", how="left")
    
    return df_base_exigida


# --- FUNCIONES DE CÁLCULO DE ICF ---

def crear_df_icf(df_a1: pd.DataFrame, df_conteo: pd.DataFrame) -> pd.DataFrame:
    """Une las frecuencias exigidas con los conteos reales y calcula el ICF."""
    df_icf = pd.merge(
        df_a1,
        df_conteo[["Fecha", "servicio", "sentido", "tipo_dia", "periodo", "expediciones_observadas"]],
        on=["Fecha", "servicio", "sentido", "tipo_dia", "periodo"],
        how="left"
    )
    df_icf = df_icf.rename(columns={"expediciones_observadas": "EO"})
    df_icf["EO"] = df_icf["EO"].fillna(0)
    df_icf['ICF'] = np.floor(np.minimum(df_icf['EE'], df_icf['EO']) / df_icf['EE'] * 100 + 0.5) / 100
    return df_icf


def round_half_up(val, decimals=2):
    multiplier = 10 ** decimals
    if isinstance(val, pd.Series):
        return np.floor(val * multiplier + 0.5) / multiplier
    else:
        return float(np.floor(val * multiplier + 0.5) / multiplier)


def calcular_psi(df, fecha_columna="Fecha", fecha_inicio_operacion=None, mas_de_24_meses=None):
    """
    Calcula el parámetro psi (ψ) según la antigüedad de la operación,
    de acuerdo a lo indicado en la Res. 49/2024 (pág. 41):
        - Hasta el mes 24 de operación: psi = 0,90
        - Desde el mes 25 en adelante:  psi = 0,95
    """
    if fecha_inicio_operacion is None and mas_de_24_meses is None:
        raise ValueError(
            "Debes especificar 'fecha_inicio_operacion' o 'mas_de_24_meses'."
        )

    if mas_de_24_meses is not None:
        return pd.Series(0.95 if mas_de_24_meses else 0.90, index=df.index)

    fecha_inicio_operacion = pd.Timestamp(fecha_inicio_operacion)

    def mes_operacion(fecha):
        return (fecha.year - fecha_inicio_operacion.year) * 12 + (fecha.month - fecha_inicio_operacion.month) + 1

    meses_operacion = df[fecha_columna].apply(mes_operacion)
    return np.where(meses_operacion <= 24, 0.90, 0.95)


def aplicar_regla_pago(icf: float, psi: float = 0.90) -> float:
    if icf < 0.50:
        return 0.50
    elif icf > psi:
        return 1.00
    else:
        return icf


def construir_resumenes_icf(df_icf: pd.DataFrame, psi_valor: float = 0.90):
    """Calcula las métricas agregadas por tipo de demanda y servicio."""
    if df_icf.empty:
        return pd.Series(dtype=float), 0.0, pd.Series(dtype=float), 0.5
        
    tabla_por_tipo_demanda = round_half_up(df_icf.groupby("tipo_demanda")["ICF"].mean(), 2)
    icf_general = round_half_up(tabla_por_tipo_demanda.mean(), 3)
    tabla_por_tipo_demanda_servicio = round_half_up(
        df_icf.groupby(["tipo_demanda", "servicio"])["ICF"].mean(), 2
    )
    
    # Aplicar la regla de pago a nivel de tipo de demanda
    tabla_por_tipo_demanda_pago = tabla_por_tipo_demanda.apply(lambda val: aplicar_regla_pago(val, psi=psi_valor))
    # Promedio de los tipos de demanda ajustados, redondeado a 2 decimales
    icf_pago = round_half_up(tabla_por_tipo_demanda_pago.mean(), 2)

    return tabla_por_tipo_demanda, icf_general, tabla_por_tipo_demanda_servicio, icf_pago


def tabla_periodo_vs_fecha(df_icf: pd.DataFrame, servicio: str, sentido: str) -> pd.DataFrame:
    """Genera la matriz periodo vs fecha con promedios."""
    df_filtro = df_icf[
        (df_icf['servicio'] == servicio) &
        (df_icf['sentido'] == sentido)
    ].copy()
    
    if df_filtro.empty:
        return pd.DataFrame()

    df_filtro['ratio_crudo'] = df_filtro['EO'] / df_filtro['EE']
    tabla = df_filtro.pivot_table(
        index='periodo',
        columns='Fecha',
        values='ratio_crudo',
        aggfunc='mean'
    )
    tabla.columns = [c.strftime('%Y-%m-%d') for c in tabla.columns]
    tabla['Promedio'] = tabla.mean(axis=1).round(2)
    return tabla


# --- PROYECCIONES ---

def proyectar_simulado_estocastico(
    df_icf_completo: pd.DataFrame,
    df_historico: pd.DataFrame,
    ultima_fecha,
    seed: int = 42
) -> pd.DataFrame:
    """
    Rellena los días faltantes del mes utilizando un muestreo estocástico con reemplazo
    del ratio de cumplimiento (EO / EE) obtenido del histórico acumulado y del mes en curso.
    """
    df_sim = df_icf_completo.copy()
    missing_mask = df_sim["Fecha"] > ultima_fecha

    # Pool de datos observados (historial + mes actual hasta la fecha)
    df_pool_actual = df_sim[~missing_mask].copy()
    if df_historico is not None and not df_historico.empty:
        # Filtrar solo columnas necesarias para evitar problemas de compatibilidad
        columnas_comunes = ["servicio", "sentido", "periodo", "tipo_dia", "tipo_demanda", "EE", "EO"]
        df_hist_filtrado = df_historico[columnas_comunes].copy()
        df_act_filtrado = df_pool_actual[columnas_comunes].copy()
        df_pool = pd.concat([df_hist_filtrado, df_act_filtrado], ignore_index=True)
    else:
        df_pool = df_pool_actual

    if df_pool.empty:
        # Fallback si no hay ningún dato: asumimos cumplimiento perfecto
        df_sim.loc[missing_mask, "EO"] = df_sim.loc[missing_mask, "EE"]
        df_sim.loc[missing_mask, "ICF"] = 1.0
        return df_sim

    # Calcular ratio de cumplimiento observado en el pool
    df_pool["ratio"] = df_pool["EO"] / df_pool["EE"]
    df_pool["ratio"] = df_pool["ratio"].fillna(0.0).clip(lower=0.0)

    # Agrupaciones en cascada para el muestreo
    pool_dict_primary = df_pool.groupby(["servicio", "sentido", "periodo", "tipo_dia"])["ratio"].apply(list).to_dict()
    pool_dict_secondary = df_pool.groupby(["tipo_demanda", "periodo"])["ratio"].apply(list).to_dict()
    pool_dict_tertiary = df_pool.groupby(["periodo"])["ratio"].apply(list).to_dict()

    # Inicializar generador aleatorio para reproducibilidad
    rng = np.random.default_rng(seed)

    df_missing = df_sim[missing_mask].copy()
    simulated_ratios = []

    for idx, row in df_missing.iterrows():
        key1 = (row["servicio"], row["sentido"], row["periodo"], row["tipo_dia"])
        key2 = (row["tipo_demanda"], row["periodo"])
        key3 = row["periodo"]

        ratio_list = pool_dict_primary.get(key1)
        if not ratio_list:
            ratio_list = pool_dict_secondary.get(key2)
        if not ratio_list:
            ratio_list = pool_dict_tertiary.get(key3)

        if ratio_list:
            ratio_val = rng.choice(ratio_list)
        else:
            ratio_val = 1.0  # Fallback absoluto

        simulated_ratios.append(ratio_val)

    # Asignar expediciones observadas simuladas (sin superar las exigidas y redondeado a entero)
    df_sim.loc[missing_mask, "EO"] = np.minimum(
        df_missing["EE"],
        np.round(df_missing["EE"] * simulated_ratios)
    )

    # Recalcular el ICF para los días proyectados
    df_sim.loc[missing_mask, "ICF"] = np.floor(
        np.minimum(df_sim.loc[missing_mask, "EE"], df_sim.loc[missing_mask, "EO"]) 
        / df_sim.loc[missing_mask, "EE"] * 100 + 0.5
    ) / 100

    return df_sim


def proyectar_ideal(df_icf_completo: pd.DataFrame, ultima_fecha) -> pd.DataFrame:
    """
    Rellena los días faltantes del mes asumiendo un escenario ideal en el cual
    se cumple exactamente la frecuencia exigida (ICF = 1.0).
    """
    df_ideal = df_icf_completo.copy()
    missing_mask = df_ideal["Fecha"] > ultima_fecha

    # En el escenario ideal, el cumplimiento de los días faltantes es perfecto
    df_ideal.loc[missing_mask, "EO"] = df_ideal.loc[missing_mask, "EE"]
    df_ideal.loc[missing_mask, "ICF"] = 1.0

    return df_ideal


# --- REPORTES COMPARATIVOS Y EXPORTACIÓN ---

def crear_tabla_comparativa(
    tabla_demanda_obs: pd.Series, icf_general_obs: float, icf_pago_obs: float,
    tabla_demanda_sim: pd.Series, icf_gen_sim: float, icf_pag_sim: float,
    tabla_demanda_ideal: pd.Series, icf_gen_ideal: float, icf_pag_ideal: float
) -> pd.DataFrame:
    """
    Construye un DataFrame comparativo side-by-side de las métricas clave.
    """
    filas = []

    # 1. Agregar ICF General
    filas.append({
        "Métrica": "ICF General (3 decimales)",
        "Real Observado (a la fecha)": icf_general_obs,
        "Proyección Simulado (Mes Completo)": icf_gen_sim,
        "Proyección Ideal (Mes Completo)": icf_gen_ideal
    })

    # 2. Agregar ICF Pago
    filas.append({
        "Métrica": "ICF Pago (Regla de Pago)",
        "Real Observado (a la fecha)": icf_pago_obs,
        "Proyección Simulado (Mes Completo)": icf_pag_sim,
        "Proyección Ideal (Mes Completo)": icf_pag_ideal
    })

    # 3. Agregar tipos de demanda
    tipos_demanda = sorted(list(
        set(tabla_demanda_obs.index) | 
        set(tabla_demanda_sim.index) | 
        set(tabla_demanda_ideal.index)
    ))

    for td in tipos_demanda:
        val_obs = tabla_demanda_obs.get(td, np.nan)
        val_sim = tabla_demanda_sim.get(td, np.nan)
        val_ideal = tabla_demanda_ideal.get(td, np.nan)

        filas.append({
            "Métrica": f"ICF Promedio Demanda: {td}",
            "Real Observado (a la fecha)": val_obs,
            "Proyección Simulado (Mes Completo)": val_sim,
            "Proyección Ideal (Mes Completo)": val_ideal
        })

    return pd.DataFrame(filas)


def escribir_tablas_resumen(ws, df_general: pd.DataFrame, df_tipo_demanda: pd.DataFrame, df_serv_demanda: pd.DataFrame) -> None:
    """
    Escribe las tablas de resumen verticalmente en una sola hoja.
    """
    font_header = Font(bold=True, size=12)
    font_table_hdr = Font(bold=True)
    
    # 1. Escribir tabla General
    ws.cell(row=2, column=1, value="ICF General").font = font_header
    for col_idx, col_name in enumerate(df_general.columns, 1):
        cell = ws.cell(row=3, column=col_idx, value=col_name)
        cell.font = font_table_hdr
    for row_idx, row_vals in enumerate(df_general.values, 4):
        for col_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            if isinstance(val, (int, float)):
                col_name_str = str(df_general.columns[col_idx-1]).lower()
                cell.number_format = '0.000' if 'general' in col_name_str else '0.00'
            
    # 2. Escribir tabla Por Tipo Demanda
    start_row_2 = 3 + len(df_general) + 3
    ws.cell(row=start_row_2 - 1, column=1, value="Por Tipo Demanda").font = font_header
    for col_idx, col_name in enumerate(df_tipo_demanda.columns, 1):
        cell = ws.cell(row=start_row_2, column=col_idx, value=col_name)
        cell.font = font_table_hdr
    for r_idx, row_vals in enumerate(df_tipo_demanda.values, start_row_2 + 1):
        for col_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(row=r_idx, column=col_idx, value=val)
            if isinstance(val, (int, float)):
                cell.number_format = '0.00'
            
    # 3. Escribir tabla Por Tipo Demanda y Servicio
    start_row_3 = start_row_2 + len(df_tipo_demanda) + 4
    ws.cell(row=start_row_3 - 1, column=1, value="Por Tipo Demanda y Servicio").font = font_header
    for col_idx, col_name in enumerate(df_serv_demanda.columns, 1):
        cell = ws.cell(row=start_row_3, column=col_idx, value=col_name)
        cell.font = font_table_hdr
    for r_idx, row_vals in enumerate(df_serv_demanda.values, start_row_3 + 1):
        for col_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(row=r_idx, column=col_idx, value=val)
            if isinstance(val, (int, float)):
                cell.number_format = '0.00'


def agregar_hoja_simulacion(
    workbook,
    df_sim: pd.DataFrame,
    sheet_name: str,
    ultima_fecha,
    label_promedio: str = "simulacion",
    label_resumen: str = "ICF simulado",
    psi_valor: float = 0.90,
) -> None:
    """
    Crea una hoja en el libro de Excel y escribe los datos día a día
    organizados en tablas paralelas (side-by-side) por tipo de demanda, incluyendo
    fórmulas de promedio de Excel, formato condicional de 3 colores y una columna
    dedicada al Pago ajustado (con fórmula IF/ROUND).
    """
    # Intentar extraer psi_valor dinámicamente si el DataFrame tiene la columna
    if "psi" in df_sim.columns and not df_sim.empty:
        psi_valor = float(df_sim["psi"].iloc[0])

    # Orden lógico de las demandas
    demand_order = ["BAJA", "MEDIA", "ALTA"]
    tipos_demanda = [td for td in demand_order if td in df_sim["tipo_demanda"].dropna().unique()]

    # Mapeo de sentidos
    sentido_map = {"I": "Ida", "R": "Reg"}

    # Bordes de distinción
    orange_side = Side(border_style="thin", color="FFA500")
    sim_border = Border(top=orange_side, bottom=orange_side, left=orange_side, right=orange_side)
    
    gray_side = Side(border_style="thin", color="D3D3D3")
    real_border = Border(top=gray_side, bottom=gray_side, left=gray_side, right=gray_side)

    align_center = Alignment(horizontal="center", vertical="center")
    
    ws = workbook.create_sheet(title=sheet_name)
    ws.views.sheetView[0].showGridLines = True

    start_col = 1
    average_cells = []
    pago_cells = []

    for td in tipos_demanda:
        # Filtrar por tipo de demanda
        df_td = df_sim[df_sim["tipo_demanda"] == td].copy()
        if df_td.empty:
            continue

        # Ordenar cronológicamente
        df_td = df_td.sort_values(["Fecha", "periodo"])

        # Obtener las combinaciones únicas de servicio/sentido para esta demanda
        combinaciones = df_td[["servicio", "sentido"]].drop_duplicates().sort_values(["servicio", "sentido"])
        if combinaciones.empty:
            continue

        # Pivotear para obtener Fecha/Periodo como índice y (servicio, sentido) como columnas
        pivot = df_td.pivot_table(
            index=["Fecha", "periodo"],
            columns=["servicio", "sentido"],
            values="ICF",
            aggfunc="first"
        )
        pivot = pivot.sort_index(level=["Fecha", "periodo"])

        # Escribir cabeceras
        ws.cell(row=2, column=start_col + 1, value="Promedio de Frecuencia").font = Font(bold=True)
        ws.cell(row=2, column=start_col + 3, value="Servicio_dicc").font = Font(bold=True)
        ws.cell(row=2, column=start_col + 4, value="Sentido").font = Font(bold=True)

        ws.cell(row=4, column=start_col, value="Dia").font = Font(bold=True)
        ws.cell(row=4, column=start_col + 1, value="Fecha").font = Font(bold=True)
        ws.cell(row=4, column=start_col + 2, value="Periodo_dicc").font = Font(bold=True)

        ws.cell(row=4, column=start_col).alignment = align_center
        ws.cell(row=4, column=start_col + 1).alignment = align_center
        ws.cell(row=4, column=start_col + 2).alignment = align_center

        data_cols = []
        for idx, (_, row_comb) in enumerate(combinaciones.iterrows()):
            srv = row_comb["servicio"]
            sent = row_comb["sentido"]
            col_idx = start_col + 3 + idx
            data_cols.append(col_idx)

            # Escribir servicio en fila 3
            ws.cell(row=3, column=col_idx, value=srv).font = Font(bold=True)
            ws.cell(row=3, column=col_idx).alignment = align_center
            # Escribir sentido en fila 4
            ws.cell(row=4, column=col_idx, value=sentido_map.get(sent, sent)).font = Font(bold=True)
            ws.cell(row=4, column=col_idx).alignment = align_center

        # Escribir filas de datos
        num_rows = len(pivot)
        for row_idx, ((fecha, per), row_vals) in enumerate(pivot.iterrows()):
            excel_row = 5 + row_idx
            fecha_val = pd.to_datetime(fecha)
            is_simulated = (fecha_val > pd.to_datetime(ultima_fecha))
            current_border = sim_border if is_simulated else real_border

            # Columna Dia (Fórmula TEXT)
            fecha_cell = f"{get_column_letter(start_col + 1)}{excel_row}"
            cell_dia = ws.cell(row=excel_row, column=start_col, value=f'=TEXT({fecha_cell},"dddd")')
            cell_dia.alignment = align_center
            cell_dia.border = current_border

            # Columna Fecha
            cell_fecha = ws.cell(row=excel_row, column=start_col + 1, value=fecha_val.date())
            cell_fecha.number_format = 'dd-mm-yy'
            cell_fecha.alignment = align_center
            cell_fecha.border = current_border

            # Columna Periodo
            cell_per = ws.cell(row=excel_row, column=start_col + 2, value=int(per))
            cell_per.alignment = align_center
            cell_per.border = current_border

            # Columnas de datos (ICF)
            for idx, (_, row_comb) in enumerate(combinaciones.iterrows()):
                srv = row_comb["servicio"]
                sent = row_comb["sentido"]
                val = row_vals.get((srv, sent))
                col_idx = start_col + 3 + idx

                cell_val = ws.cell(row=excel_row, column=col_idx)
                cell_val.alignment = align_center
                cell_val.border = current_border

                if pd.notna(val):
                    cell_val.value = float(val)
                    cell_val.number_format = '0.00'

        # Escribir el promedio de este tipo de demanda
        end_col = start_col + 2 + len(combinaciones)
        avg_col = end_col + 2
        pago_col = end_col + 3

        # 1. Promedio simple
        ws.cell(row=3, column=avg_col, value=f'Promedio {label_promedio} "{td}"').font = Font(bold=True)
        ws.cell(row=3, column=avg_col).alignment = align_center

        first_data_cell = f"{get_column_letter(start_col + 3)}5"
        last_data_cell = f"{get_column_letter(end_col)}{5 + num_rows - 1}"
        ws.cell(row=4, column=avg_col, value=f'=ROUND(AVERAGE({first_data_cell}:{last_data_cell}), 2)').font = Font(bold=True)
        ws.cell(row=4, column=avg_col).alignment = align_center
        ws.cell(row=4, column=avg_col).number_format = '0.00000'

        average_cells.append(f"{get_column_letter(avg_col)}4")

        # 2. Promedio Pago (regla de pago 0.5 y psi_valor a nivel de demanda)
        ws.cell(row=3, column=pago_col, value=f'Pago {label_promedio} "{td}"').font = Font(bold=True)
        ws.cell(row=3, column=pago_col).alignment = align_center

        avg_cell_ref = f"{get_column_letter(avg_col)}4"
        pago_formula = f"=IF({avg_cell_ref}<0.5, 0.5, IF({avg_cell_ref}>{psi_valor:.2f}, 1.0, {avg_cell_ref}))"
        ws.cell(row=4, column=pago_col, value=pago_formula).font = Font(bold=True)
        ws.cell(row=4, column=pago_col).alignment = align_center
        ws.cell(row=4, column=pago_col).number_format = '0.00000'

        pago_cells.append(f"{get_column_letter(pago_col)}4")

        # Formato condicional (escala de 3 colores) para el rango de datos
        cell_range = f"{get_column_letter(start_col + 3)}5:{get_column_letter(end_col)}{5 + num_rows - 1}"
        color_scale = ColorScaleRule(
            start_type='min', start_color='FFF8696B', # Rojo suave
            mid_type='percentile', mid_value=50, mid_color='FFFFEB84', # Amarillo suave
            end_type='max', end_color='FF63BE7B' # Verde suave
        )
        ws.conditional_formatting.add(cell_range, color_scale)

        # Siguiente tabla empieza 5 columnas después (deja 1 columna en blanco tras el Pago)
        start_col = end_col + 5

    # Escribir resumen general al final
    if average_cells:
        summary_label_col = start_col
        summary_val_col = start_col + 1

        # ICF Promedio
        ws.cell(row=3, column=summary_label_col, value=label_resumen).font = Font(bold=True)
        ws.cell(row=3, column=summary_label_col).alignment = align_center

        avg_formula_terms = "+".join(average_cells)
        formula_avg = f"=({avg_formula_terms})/{len(average_cells)}"
        cell_summary = ws.cell(row=3, column=summary_val_col, value=formula_avg)
        cell_summary.font = Font(bold=True)
        cell_summary.alignment = align_center
        cell_summary.number_format = '0.00000'

        # ICF Pago
        ws.cell(row=5, column=summary_label_col, value=f"{label_resumen} Pago").font = Font(bold=True)
        ws.cell(row=5, column=summary_label_col).alignment = align_center

        pago_formula_terms = "+".join(pago_cells)
        formula_pago = f"=ROUND(({pago_formula_terms})/{len(pago_cells)}, 2)"
        cell_pago_summary = ws.cell(row=5, column=summary_val_col, value=formula_pago)
        cell_pago_summary.font = Font(bold=True)
        cell_pago_summary.alignment = align_center
        cell_pago_summary.number_format = '0.00000'


def exportar_resumenes_icf(
    tabla_por_tipo_demanda: pd.Series,
    icf_general: float,
    tabla_por_tipo_demanda_servicio: pd.Series,
    icf_pago: float,
    ruta_archivo,
    df_icf: pd.DataFrame = None,
) -> None:
    """
    Exporta los resúmenes del ICF a un único Excel.
    La primera pestaña es "Detalle_Diario" (si se entrega df_icf) y la
    segunda es "Resumen", que agrupa todas las tablas de resumen verticalmente.
    """
    wb = Workbook()

    # 1. Crear 'Detalle_Diario' si df_icf existe
    if df_icf is not None and not df_icf.empty:
        ultima_fecha_real = df_icf["Fecha"].max()
        agregar_hoja_simulacion(
            wb,
            df_icf,
            "Detalle_Diario",
            ultima_fecha_real,
            label_promedio="observado",
            label_resumen="ICF observado",
        )
        if "Sheet" in wb.sheetnames:
            wb.remove(wb["Sheet"])

    # 2. Crear 'Resumen'
    ws_resumen = wb.create_sheet(title="Resumen")
    ws_resumen.views.sheetView[0].showGridLines = True

    df_tipo_demanda = tabla_por_tipo_demanda.reset_index().rename(
        columns={"ICF": "ICF_promedio"}
    )
    
    psi_valor = 0.90
    if df_icf is not None and not df_icf.empty and "psi" in df_icf.columns:
        psi_valor = float(df_icf["psi"].iloc[0])

    df_tipo_demanda["ICF_pago"] = df_tipo_demanda["ICF_promedio"].apply(
        lambda val: aplicar_regla_pago(val, psi=psi_valor)
    )
    df_serv_demanda = tabla_por_tipo_demanda_servicio.reset_index().rename(
        columns={"ICF": "ICF_promedio"}
    )
    df_general = pd.DataFrame({"ICF_general": [icf_general], "ICF_pago": [icf_pago]})

    escribir_tablas_resumen(ws_resumen, df_general, df_tipo_demanda, df_serv_demanda)

    # Si 'Sheet' por defecto sigue ahí, la removemos
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])

    wb.save(ruta_archivo)


def exportar_reporte_proyeccion(
    df_comparativa: pd.DataFrame,
    df_sim: pd.DataFrame,
    df_ideal: pd.DataFrame,
    tabla_demanda_sim: pd.Series, icf_gen_sim: float, tabla_serv_sim: pd.Series, icf_pag_sim: float,
    tabla_demanda_ideal: pd.Series, icf_gen_ideal: float, tabla_serv_ideal: pd.Series, icf_pag_ideal: float,
    ultima_fecha,
    ruta_archivo
) -> None:
    """
    Exporta todos los resultados de simulación y proyecciones a un archivo Excel.
    """
    wb = Workbook()

    # 1. Agregar 'Simulacion' (primera pestaña)
    agregar_hoja_simulacion(wb, df_sim, "Simulacion", ultima_fecha)
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])

    # 2. Agregar 'Simulacion_Ideal' (segunda pestaña)
    agregar_hoja_simulacion(wb, df_ideal, "Simulacion_Ideal", ultima_fecha)

    # 3. Agregar 'Comparativa' (tercera pestaña)
    ws_comp = wb.create_sheet(title="Comparativa")
    ws_comp.views.sheetView[0].showGridLines = True
    for col_idx, col_name in enumerate(df_comparativa.columns, 1):
        ws_comp.cell(row=1, column=col_idx, value=col_name).font = Font(bold=True)
    for r_idx, row_vals in enumerate(df_comparativa.values, 2):
        for col_idx, val in enumerate(row_vals, 1):
            ws_comp.cell(row=r_idx, column=col_idx, value=val)

    # 4. Agregar 'Resumen_Simulado' (cuarta pestaña)
    ws_res_sim = wb.create_sheet(title="Resumen_Simulado")
    ws_res_sim.views.sheetView[0].showGridLines = True
    df_demanda_sim = tabla_demanda_sim.reset_index().rename(columns={"ICF": "ICF_promedio"})
    
    psi_valor_sim = 0.90
    if "psi" in df_sim.columns and not df_sim.empty:
        psi_valor_sim = float(df_sim["psi"].iloc[0])
        
    df_demanda_sim["ICF_pago"] = df_demanda_sim["ICF_promedio"].apply(lambda val: aplicar_regla_pago(val, psi=psi_valor_sim))
    df_serv_sim = tabla_serv_sim.reset_index().rename(columns={"ICF": "ICF_promedio"})
    df_gen_sim = pd.DataFrame({"ICF_general": [icf_gen_sim], "ICF_pago": [icf_pag_sim]})
    escribir_tablas_resumen(ws_res_sim, df_gen_sim, df_demanda_sim, df_serv_sim)

    # 5. Agregar 'Resumen_Ideal' (quinta pestaña)
    ws_res_ideal = wb.create_sheet(title="Resumen_Ideal")
    ws_res_ideal.views.sheetView[0].showGridLines = True
    df_demanda_ideal = tabla_demanda_ideal.reset_index().rename(columns={"ICF": "ICF_promedio"})
    
    psi_valor_ideal = 0.90
    if "psi" in df_ideal.columns and not df_ideal.empty:
        psi_valor_ideal = float(df_ideal["psi"].iloc[0])
        
    df_demanda_ideal["ICF_pago"] = df_demanda_ideal["ICF_promedio"].apply(lambda val: aplicar_regla_pago(val, psi=psi_valor_ideal))
    df_serv_ideal = tabla_serv_ideal.reset_index().rename(columns={"ICF": "ICF_promedio"})
    df_gen_ideal = pd.DataFrame({"ICF_general": [icf_gen_ideal], "ICF_pago": [icf_pag_ideal]})
    escribir_tablas_resumen(ws_res_ideal, df_gen_ideal, df_demanda_ideal, df_serv_ideal)

    wb.save(ruta_archivo)


# --- ORQUESTADOR PRINCIPAL DEL CÁLCULO ---

def ejecutar_calculo_icf(operador: str, anio: int, mes: int, db_path: str) -> Tuple[bytes, bytes, str]:
    """
    Ejecuta el cálculo completo del ICF y proyecciones para un operador, año y mes.
    Retorna (reporte_excel_bytes, proyeccion_excel_bytes, resumen_texto_whatsapp).
    """
    hoy = datetime.now()
    es_mes_en_curso = (anio == hoy.year and mes == hoy.month)
    es_mes_futuro = (anio > hoy.year) or (anio == hoy.year and mes > hoy.month)
    
    # 1. Obtener conteos de expediciones reales del mes actual
    df_conteo = obtener_df_conteo_db(db_path, operador, anio, mes)
    
    fecha_corte = None
    if es_mes_en_curso:
        if df_conteo.empty:
            fecha_corte = pd.Timestamp(hoy.date())
        else:
            fecha_corte = df_conteo["Fecha"].max()
    elif es_mes_futuro:
        fecha_corte = pd.Timestamp(hoy.date())
        
    # Calendario observado hasta la fecha de corte
    calendario = construir_calendario_mes(anio, mes, fecha_corte=fecha_corte)
    
    # 2. Cargar frecuencias exigidas cruzadas con el calendario
    df_a1 = obtener_df_a1_db(db_path, operador, calendario)
    if df_a1.empty:
        raise ValueError(f"No hay frecuencias programadas (Anexo 1) para el operador '{operador}' en la base de datos.")
        
    # 3. Crear df_icf observado (a la fecha)
    df_icf_obs = crear_df_icf(df_a1, df_conteo)
    
    # Calcular psi dinámicamente (por defecto mas_de_24_meses=False para usar 0.90 en el bot)
    df_icf_obs["psi"] = calcular_psi(df_icf_obs, mas_de_24_meses=False)
    psi_valor = 0.90
    if not df_icf_obs.empty and "psi" in df_icf_obs.columns:
        psi_valor = float(df_icf_obs["psi"].iloc[0])
        
    # 4. Cargar histórico acumulado de meses anteriores
    inicio_mes = pd.Timestamp(year=anio, month=mes, day=1)
    df_conteo_hist = obtener_df_conteo_historico_db(db_path, operador, inicio_mes)
    if not df_conteo_hist.empty:
        fechas_hist = pd.Series(df_conteo_hist["Fecha"].unique())
        calendario_hist = pd.DataFrame({"Fecha": fechas_hist, "tipo_dia": _calcular_tipo_dia(fechas_hist)})
        df_a1_hist = obtener_df_a1_db(db_path, operador, calendario_hist)
        df_icf_hist = crear_df_icf(df_a1_hist, df_conteo_hist)
    else:
        df_icf_hist = pd.DataFrame()
        
    # 5. Proyecciones (mes completo completo)
    calendario_completo = construir_calendario_mes(anio, mes)
    df_a1_completo = obtener_df_a1_db(db_path, operador, calendario_completo)
    df_icf_completo = crear_df_icf(df_a1_completo, df_conteo)
    
    # Última fecha real observada con datos
    ultima_fecha_real = df_conteo["Fecha"].max() if not df_conteo.empty else pd.Timestamp(year=anio, month=mes, day=1) - pd.Timedelta(days=1)
    
    # Simulación Estocástica y Simulación Ideal
    df_sim = proyectar_simulado_estocastico(df_icf_completo, df_icf_hist, ultima_fecha_real)
    df_ideal = proyectar_ideal(df_icf_completo, ultima_fecha_real)
    
    # Asignar psi a los dataframes de proyecciones
    df_sim["psi"] = calcular_psi(df_sim, mas_de_24_meses=False)
    df_ideal["psi"] = calcular_psi(df_ideal, mas_de_24_meses=False)
    
    # 6. Calcular resúmenes para cada escenario
    res_td_obs, res_gen_obs, res_serv_obs, res_pago_obs = construir_resumenes_icf(df_icf_obs, psi_valor)
    res_td_sim, res_gen_sim, res_serv_sim, res_pago_sim = construir_resumenes_icf(df_sim, psi_valor)
    res_td_ideal, res_gen_ideal, res_serv_ideal, res_pago_ideal = construir_resumenes_icf(df_ideal, psi_valor)
    
    # 7. Crear tabla comparativa
    df_comparativa = crear_tabla_comparativa(
        res_td_obs, res_gen_obs, res_pago_obs,
        res_td_sim, res_gen_sim, res_pago_sim,
        res_td_ideal, res_gen_ideal, res_pago_ideal
    )
    
    # 8. Exportar reporte.xlsx (observado)
    reporte_io = io.BytesIO()
    exportar_resumenes_icf(res_td_obs, res_gen_obs, res_serv_obs, res_pago_obs, reporte_io, df_icf_obs)
    reporte_bytes = reporte_io.getvalue()
    
    # 9. Exportar reporte_proyeccion.xlsx (simulado e ideal)
    proyeccion_io = io.BytesIO()
    exportar_reporte_proyeccion(
        df_comparativa,
        df_sim,
        df_ideal,
        res_td_sim, res_gen_sim, res_serv_sim, res_pago_sim,
        res_td_ideal, res_gen_ideal, res_serv_ideal, res_pago_ideal,
        ultima_fecha_real,
        proyeccion_io
    )
    proyeccion_bytes = proyeccion_io.getvalue()
    
    # 10. Resumen texto
    resumen_txt = (
        f"📊 *Reporte ICF: {operador.upper()} - {mes:02d}/{anio}*\n\n"
        f"🟢 *Real Observado (a la fecha)*:\n"
        f"   - ICF General: *{res_gen_obs:.3f}*\n"
        f"   - ICF Pago: *{res_pago_obs:.2f}*\n\n"
        f"🔵 *Proyección Simulación Estocástica*:\n"
        f"   - ICF General: *{res_gen_sim:.3f}*\n"
        f"   - ICF Pago: *{res_pago_sim:.2f}*\n\n"
        f"⚫ *Proyección Escenario Ideal*:\n"
        f"   - ICF General: *{res_gen_ideal:.3f}*\n"
        f"   - ICF Pago: *{res_pago_ideal:.2f}*\n"
    )
    
    return reporte_bytes, proyeccion_bytes, resumen_txt
